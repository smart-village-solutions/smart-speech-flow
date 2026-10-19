#!/usr/bin/env python3
"""Create the designer handoff: slide text in DOCX and all artwork in ZIP.

Requires python-docx. Run: python docs/user-guide/build_designer_package.py
"""

from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from build_pdf import MEDIA, PAGES, ROOT, SCREENSHOTS, capture, parse_page


DOCX = ROOT / "kasseldialog-folientexte.docx"
ZIP = ROOT / "kasseldialog-grafikerpaket.zip"
GRAPHICS = ROOT / "graphics"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))")


def add_hyperlink(paragraph, label: str, target: str, *, bold: bool = False) -> None:
    relationship = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    if bold:
        properties.append(OxmlElement("w:b"))
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = label
    run.append(value)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_inline(paragraph, value: str, *, bold: bool = False) -> None:
    for part in INLINE.split(value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_inline(paragraph, part[2:-2], bold=True)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, target = part[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, target, bold=bold)
        else:
            run = paragraph.add_run(part)
            run.bold = bold


def make_docx(paths: list[Path]) -> None:
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(1.7)
    section.left_margin = section.right_margin = Cm(2.1)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(24, 50, 71)
    normal.paragraph_format.space_after = Pt(10)
    heading = document.styles["Heading 1"]
    heading.font.name = "Aptos Display"
    heading.font.size = Pt(23)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(29, 99, 157)
    heading.paragraph_format.space_after = Pt(18)
    for number, path in enumerate(paths, start=1):
        if number > 1:
            document.add_page_break()
        title, blocks = parse_page(path)
        document.add_heading(f"{number:02d} · {title}", level=1)
        for kind, content in blocks:
            if kind == "table":
                rows = content
                table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
                table.style = "Table Grid"
                for i, row in enumerate(rows):
                    for j, cell in enumerate(row):
                        paragraph = table.cell(i, j).paragraphs[0]
                        add_inline(paragraph, cell, bold=i == 0)
                document.add_paragraph()
            else:
                paragraph = document.add_paragraph()
                if kind == "step":
                    paragraph.paragraph_format.left_indent = Cm(0.45)
                add_inline(paragraph, str(content))
    footer = section.footer.paragraphs[0]
    footer.text = "KasselDIALOG · Folientexte für die Gestaltung · Arbeitsfassung"
    footer.style = document.styles["Caption"]
    document.save(DOCX)


def slide_artwork() -> dict[int, list[str]]:
    used: dict[int, list[str]] = {
        2: ["grafiken/G02.svg (auch als PDF)"],
        3: ["grafiken/G03.svg (auch als PDF)"],
        17: ["grafiken/G04.svg (auch als PDF)"],
    }
    for number, prefixes in {4: ["P00", "P01", "P02"], 5: ["P03", "S01", "S02"]}.items():
        used[number] = [f"screenshots/{capture(prefix).name}" for prefix in prefixes]
    for number, items in MEDIA.items():
        used[number] = [f"screenshots/{capture(prefix, index).name}" for _, _, prefix, index in items if prefix and capture(prefix, index)]
    return used


def make_readme(paths: list[Path]) -> str:
    used = slide_artwork()
    included = {Path(name).name for files in used.values() for name in files}
    extras = sorted(path.name for path in SCREENSHOTS.glob("*.png") if path.name not in included)
    lines = [
        "# KasselDIALOG · Übergabe an die Gestaltung",
        "",
        "Arbeitsfassung. Das gesondert bereitgestellte PDF zeigt die aktuelle Anordnung.",
        "`kasseldialog-folientexte.docx` enthält die Texte aller 26 Folien, ohne eingebettete Bilder.",
        "`screenshots/` enthält alle vorhandenen Originalaufnahmen; `grafiken/` enthält G02–G04 als SVG und Vektor-PDF.",
        "",
        "## Bildzuordnung im aktuellen PDF",
        "",
    ]
    for number, path in enumerate(paths, start=1):
        title, _ = parse_page(path)
        art = ", ".join(used.get(number, [])) or "kein Bild"
        lines.append(f"- Folie {number:02d} · {title}: {art}")
    lines.extend(["", "## Zusätzlich enthaltene Originalaufnahmen", ""])
    lines.extend(f"- screenshots/{name}" for name in extras)
    lines.extend([
        "",
        "## Vor Veröffentlichung prüfen",
        "",
        "- Die auf S05 und S04 sichtbaren Sitzungscodes, Einladungslinks und QR-Ziele dürfen nicht mehr nutzbar sein.",
        "- Produktversion, Datenschutzinformation und Wortlaut der Ansichten müssen fachlich freigegeben werden.",
        "- S09 ist derzeit eine Gästeansicht. Das PDF verwendet für Folie 09 die Mitarbeitendenansicht S13a.",
        "- G05 und G06 wurden als Bilder gestrichen; die zugehörigen Folientexte bleiben erhalten.",
        "",
    ])
    return "\n".join(lines)


def main() -> None:
    paths = sorted(PAGES.glob("[0-9][0-9]-*.md"))
    if len(paths) != 26:
        raise ValueError(f"Expected 26 slide files, found {len(paths)}")
    make_docx(paths)
    with ZipFile(ZIP, "w", compression=ZIP_DEFLATED) as archive:
        archive.write(DOCX, DOCX.name)
        archive.writestr("README.md", make_readme(paths))
        for path in sorted(SCREENSHOTS.glob("*.png")):
            archive.write(path, f"screenshots/{path.name}")
        for path in sorted(GRAPHICS.glob("G0[234].*")):
            archive.write(path, f"grafiken/{path.name}")
    print(f"Created {DOCX} and {ZIP}")


if __name__ == "__main__":
    main()
